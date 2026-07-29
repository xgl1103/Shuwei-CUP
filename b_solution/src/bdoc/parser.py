from __future__ import annotations

import logging
import re
import io
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable, Mapping, Protocol, runtime_checkable
from xml.etree import ElementTree as ET

import numpy as np
from PIL import Image
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("pypdf._reader").setLevel(logging.ERROR)

OcrBackend = Callable[[Path | Image.Image], str]


@runtime_checkable
class OcrBackendProtocol(Protocol):
    def __call__(self, source: Path | Image.Image) -> str: ...


@dataclass(slots=True)
class ParseConfig:
    max_pdf_pages: int = 25
    max_xlsx_rows: int = 40
    max_xlsx_cols: int = 15
    include_hidden_sheets: bool = True
    include_cell_comments: bool = True
    max_chars: int = 20_000
    include_ocr: bool = True
    ocr_language: str = "chi_sim+eng"
    pdf_ocr_scale: float = 2.0
    docx_ocr_min_text_chars: int = 0
    pdf_ocr_min_text_chars: int = 0


@dataclass(slots=True)
class ParsedDocument:
    doc_id: str
    dataset: str
    file_path: str
    file_name: str
    suffix: str
    size_bytes: int
    title: str
    time_info: str
    raw_text: str
    clean_text: str
    table_text: str
    ocr_text: str
    parse_status: str
    parse_notes: str
    page_count: int | None
    sheet_count: int | None
    image_width: int | None
    image_height: int | None
    content_truncated: bool
    char_count: int


def to_dict(record: ParsedDocument) -> dict[str, object]:
    return asdict(record)


def strip_surrogates(text: str) -> str:
    if not text:
        return ""
    return "".join(ch for ch in text if not 0xD800 <= ord(ch) <= 0xDFFF)


def normalize_text(text: str) -> str:
    if not text:
        return ""
    text = strip_surrogates(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u3000", " ")
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def truncate_text(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


def guess_title(text: str, fallback: str) -> str:
    text = strip_surrogates(text)
    lines = [line.strip(" \t-—:：|") for line in text.splitlines() if line.strip()]
    for line in lines[:10]:
        if 4 <= len(line) <= 120 and not re.fullmatch(r"[\W_0-9]+", line):
            return line
    return fallback


def extract_time_info(text: str, max_matches: int = 4) -> str:
    text = strip_surrogates(text)
    patterns = [
        r"20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}日?",
        r"20\d{2}[年/-]\d{1,2}[月/-]?",
        r"20\d{2}-\d{2}",
        r"20\d{2}Q[1-4]",
        r"20\d{2}年Q[1-4]",
    ]
    matches: list[str] = []
    seen: set[str] = set()
    for pattern in patterns:
        for match in re.finditer(pattern, text):
            value = match.group(0)
            if value not in seen:
                seen.add(value)
                matches.append(value)
                if len(matches) >= max_matches:
                    return "; ".join(matches)
    return "; ".join(matches)


def load_text_file(path: Path) -> tuple[str, str]:
    encodings = ("utf-8-sig", "utf-8", "gb18030", "gbk")
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding), f"encoding={encoding}"
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            return path.read_text(encoding=encoding, errors="ignore"), f"encoding={encoding}; warning={exc}"
    return path.read_text(encoding="utf-8", errors="ignore"), "encoding=utf-8; warning=lossy"


def extract_docx_media_ocr(
    path: Path,
    ocr_backend: OcrBackendProtocol | None,
    max_images: int = 40,
) -> tuple[str, str]:
    if ocr_backend is None:
        return "", "docx_image_ocr_unavailable"

    notes: list[str] = []
    chunks: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            media_names = [
                name
                for name in archive.namelist()
                if name.startswith("word/media/")
                and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
            ]
            if len(media_names) > max_images:
                notes.append(f"docx_image_ocr_truncated={len(media_names)}->{max_images}")
            for name in media_names[:max_images]:
                try:
                    with Image.open(io.BytesIO(archive.read(name))) as image:
                        text = _ocr_source(ocr_backend, image.convert("RGB"))
                    if text.strip():
                        chunks.append(text)
                except Exception as exc:
                    notes.append(f"docx_image_ocr_error={Path(name).name}:{exc}")
    except Exception as exc:
        return "", f"docx_image_ocr_open_error={exc}"

    if chunks:
        notes.append(f"docx_image_ocr_used={len(chunks)}")
    elif not notes:
        notes.append("docx_image_ocr_empty")
    return "\n\n".join(chunks), "; ".join(notes)


def _strip_xml_tag_text(value: str) -> str:
    value = re.sub(r"<[^>]+>", "", value)
    value = (
        value.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&amp;", "&")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
    )
    return normalize_text(value)


def extract_docx_extra_xml_text(path: Path) -> tuple[str, str]:
    notes: list[str] = []
    chunks: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            xml_names = [
                name
                for name in archive.namelist()
                if name.startswith("word/")
                and name.endswith(".xml")
                and (
                    name.startswith("word/header")
                    or name.startswith("word/footer")
                    or name in {"word/document.xml", "word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
                )
            ]
            namespaces = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                "v": "urn:schemas-microsoft-com:vml",
            }
            seen: set[str] = set()
            for name in xml_names:
                xml = archive.read(name)
                try:
                    root = ET.fromstring(xml)
                except Exception:
                    raw = xml.decode("utf-8", errors="ignore")
                    values = re.findall(r"<(?:w:t|a:t)[^>]*>(.*?)</(?:w:t|a:t)>", raw)
                    values += re.findall(r"<v:textbox[^>]*>(.*?)</v:textbox>", raw, flags=re.S)
                    for value in values:
                        text = _strip_xml_tag_text(value)
                        if text and text not in seen:
                            seen.add(text)
                            chunks.append(text)
                    continue

                values: list[str] = []
                for tag in (".//w:t", ".//a:t"):
                    values.extend(elem.text or "" for elem in root.findall(tag, namespaces))
                for elem in root.iter():
                    if elem.text and elem.tag.endswith("}t"):
                        values.append(elem.text)
                for value in values:
                    text = normalize_text(value)
                    if text and text not in seen:
                        seen.add(text)
                        chunks.append(text)
            if chunks:
                notes.append(f"docx_extra_xml_text={len(chunks)}")
    except Exception as exc:
        notes.append(f"docx_extra_xml_error={exc}")
    return "\n".join(chunks), "; ".join(notes)


def extract_docx_text(
    path: Path,
    include_ocr: bool = False,
    ocr_backend: OcrBackendProtocol | None = None,
    ocr_min_text_chars: int = 0,
) -> tuple[str, str, str, str]:
    notes: list[str] = []
    document = Document(str(path))

    title = document.core_properties.title or ""
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]

    header_footer_texts: list[str] = []
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            for paragraph in container.paragraphs:
                text = normalize_text(paragraph.text)
                if text:
                    header_footer_texts.append(text)
            for table in container.tables:
                for row in table.rows:
                    cells = [normalize_text(cell.text.replace("\n", " ")) for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        header_footer_texts.append(line)

    table_blocks: list[str] = []
    for index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            table_blocks.append(f"[Table {index}]\n" + "\n".join(rows))

    text = "\n".join(paragraphs)
    table_text = "\n\n".join(table_blocks)
    extra_xml_text, extra_note = extract_docx_extra_xml_text(path)
    extra_parts = [part for part in ("\n".join(header_footer_texts), extra_xml_text) if part]
    if extra_parts:
        text = "\n".join(part for part in (text, "\n".join(extra_parts)) if part)
    if extra_note:
        notes.append(extra_note)
    text_for_ocr_decision = normalize_text("\n".join(part for part in (text, table_text) if part))
    should_ocr_images = include_ocr and len(text_for_ocr_decision) <= max(0, ocr_min_text_chars)
    if should_ocr_images:
        ocr_text, ocr_note = extract_docx_media_ocr(path, ocr_backend)
        if ocr_text.strip():
            text = "\n\n".join(part for part in (text, "[Embedded Image OCR]", ocr_text) if part)
        notes.append(ocr_note)
    if not title and paragraphs:
        title = paragraphs[0]
    return title, text, table_text, "; ".join(notes)


def _ocr_source(ocr_backend: OcrBackendProtocol | None, source: Path | Image.Image) -> str:
    if ocr_backend is None:
        return ""
    return ocr_backend(source)


def extract_pdf_text(
    path: Path,
    max_pages: int,
    include_ocr: bool = False,
    ocr_backend: OcrBackendProtocol | None = None,
    ocr_scale: float = 2.0,
    ocr_min_text_chars: int = 0,
) -> tuple[str, str, str, str, int, str]:
    reader = PdfReader(str(path), strict=False)
    page_count = len(reader.pages)
    title = ""
    try:
        metadata = reader.metadata
        if metadata and metadata.title:
            title = str(metadata.title).strip()
    except Exception:
        pass

    pages_to_read = min(page_count, max_pages)
    if page_count > max_pages:
        notes = [f"page_truncated={page_count}->{max_pages}"]
    else:
        notes = []

    chunks: list[str] = []
    ocr_chunks: list[str] = []
    matrix = None
    pdf_module = None
    if include_ocr and ocr_backend is not None:
        try:
            import fitz  # type: ignore
        except Exception:
            pdf_module = None
        else:
            pdf_module = fitz
            matrix = fitz.Matrix(ocr_scale, ocr_scale)

    for index in range(pages_to_read):
        try:
            page_text = reader.pages[index].extract_text() or ""
        except Exception as exc:
            page_text = ""
            notes.append(f"page_{index + 1}_error={exc}")
        if page_text.strip():
            chunks.append(page_text)

    extracted_text_chars = len(normalize_text("\n\n".join(chunks)))
    should_ocr_pdf = include_ocr and ocr_backend is not None and (
        not chunks or extracted_text_chars <= max(0, ocr_min_text_chars)
    )
    if should_ocr_pdf:
        if pdf_module is None or matrix is None:
            notes.append("pdf_ocr_unavailable")
        else:
            try:
                pdf_doc = pdf_module.open(str(path))
            except Exception as exc:
                notes.append(f"pdf_ocr_open_error={exc}")
            else:
                try:
                    pages_to_ocr = min(len(pdf_doc), max_pages)
                    for index in range(pages_to_ocr):
                        try:
                            page = pdf_doc.load_page(index)
                            pix = page.get_pixmap(matrix=matrix, alpha=False)
                            image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                            text = _ocr_source(ocr_backend, image)
                            if text.strip():
                                ocr_chunks.append(text)
                        except Exception as exc:
                            notes.append(f"pdf_ocr_page_{index + 1}_error={exc}")
                finally:
                    pdf_doc.close()
        if ocr_chunks:
            notes.append(f"pdf_ocr_pages={len(ocr_chunks)}")
        else:
            notes.append("pdf_ocr_empty")

    return (
        title,
        "\n\n".join(chunks),
        "",
        "\n\n".join(ocr_chunks),
        page_count,
        "; ".join(notes),
    )


def extract_xlsx_text(
    path: Path,
    max_rows: int,
    max_cols: int,
    include_hidden_sheets: bool = True,
    include_cell_comments: bool = True,
) -> tuple[str, str, str, int, str]:
    workbook = load_workbook(path, read_only=not include_cell_comments, data_only=True)
    notes: list[str] = []
    sheet_count = len(workbook.sheetnames)
    title = workbook.properties.title or ""
    summary_blocks: list[str] = []
    table_blocks: list[str] = []

    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible" and not include_hidden_sheets:
            notes.append(f"hidden_sheet_skipped={sheet.title}")
            continue
        rows: list[str] = []
        row_count = 0
        comment_blocks: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            row_count += 1
            if row_count > max_rows:
                break
            values = ["" if value is None else str(value).strip() for value in row[:max_cols]]
            if any(values):
                rows.append(" | ".join(values))
        if include_cell_comments:
            try:
                for row in sheet.iter_rows():
                    for cell in row:
                        if getattr(cell, "comment", None) and cell.comment.text:
                            comment_blocks.append(f"{cell.coordinate}: {normalize_text(cell.comment.text)}")
            except Exception as exc:
                notes.append(f"comment_extract_error={sheet.title}:{exc}")
        summary_blocks.append(f"[Sheet {sheet.title} rows={sheet.max_row} cols={sheet.max_column}]")
        if rows:
            table_blocks.append("\n".join(rows))
        if comment_blocks:
            table_blocks.append(f"[Comments {sheet.title}]\n" + "\n".join(comment_blocks[: max_rows]))
            notes.append(f"comments={sheet.title}:{len(comment_blocks)}")

    if not title:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=min(5, sheet.max_row), values_only=True):
                for value in row:
                    if value not in (None, ""):
                        title = str(value).strip()
                        break
                if title:
                    break
            if title:
                break

    workbook.close()
    return (
        title,
        "\n".join(summary_blocks),
        "\n\n".join(table_blocks),
        sheet_count,
        "; ".join(notes),
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return normalize_text(str(value))


def _merge_time_info(*values: str) -> str:
    seen: set[str] = set()
    merged: list[str] = []
    for value in values:
        for item in re.split(r"[;；]\s*", value or ""):
            item = item.strip()
            if item and item not in seen:
                seen.add(item)
                merged.append(item)
    return "; ".join(merged)


def _dataset3_title(file_id: str, body: str) -> str:
    snippet = normalize_text(body).replace("\n", " ")
    if len(snippet) > 64:
        snippet = snippet[:64] + "..."
    return strip_surrogates(f"{file_id} {snippet}".strip() or file_id)


def parse_dataset3_xlsx_rows(row: Mapping[str, str], config: ParseConfig) -> list[ParsedDocument]:
    file_path = Path(row["file_path"])
    dataset = row.get("dataset") or "dataset3"
    file_name = row.get("file_name") or file_path.name
    suffix = (row.get("suffix") or file_path.suffix).lower()
    size_bytes = int(row.get("size_bytes") or file_path.stat().st_size)

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    records: list[ParsedDocument] = []
    try:
        sheet = workbook.worksheets[0]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [_cell_text(value) for value in header_values]
        header_map = {name: index for index, name in enumerate(headers) if name}

        id_index = header_map.get("文件编号", 0)
        body_index = header_map.get("正文片段", 1)
        time_index = header_map.get("时间信息", 2)

        for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            file_id = _cell_text(values[id_index]) if id_index < len(values) else ""
            body = _cell_text(values[body_index]) if body_index < len(values) else ""
            source_time = _cell_text(values[time_index]) if time_index < len(values) else ""
            if not file_id and not body and not source_time:
                continue
            if not file_id:
                file_id = f"row{row_index:05d}"

            field_text = normalize_text(
                "\n".join(
                    part
                    for part in (
                        f"文件编号：{file_id}",
                        f"时间信息：{source_time}" if source_time else "",
                        "正文片段：",
                        body,
                    )
                    if part
                )
            )
            clean_text, truncated = truncate_text(field_text, config.max_chars)
            title = _dataset3_title(file_id, body)
            time_info = _merge_time_info(source_time, extract_time_info(body))
            parse_status = "ok" if body else "empty"
            notes = (
                f"dataset3_row={row_index}; "
                f"source_xlsx={file_name}; "
                f"columns={','.join(headers)}"
            )

            records.append(
                ParsedDocument(
                    doc_id=f"{dataset}:{file_id}",
                    dataset=dataset,
                    file_path=str(file_path),
                    file_name=file_name,
                    suffix=suffix,
                    size_bytes=size_bytes,
                    title=title,
                    time_info=strip_surrogates(time_info),
                    raw_text=strip_surrogates(clean_text),
                    clean_text=strip_surrogates(clean_text),
                    table_text="",
                    ocr_text="",
                    parse_status=parse_status,
                    parse_notes=strip_surrogates(notes),
                    page_count=None,
                    sheet_count=1,
                    image_width=None,
                    image_height=None,
                    content_truncated=truncated,
                    char_count=len(clean_text),
                )
            )
    finally:
        workbook.close()

    return records


def build_ocr_backend(language: str) -> tuple[OcrBackendProtocol | None, str]:
    try:
        import pytesseract  # type: ignore
    except Exception:
        pytesseract = None

    if pytesseract is not None:
        def pytesseract_backend(source: Path | Image.Image) -> str:
            if isinstance(source, Path):
                with Image.open(source) as image:
                    image = image.convert("RGB")
                    try:
                        return pytesseract.image_to_string(image, lang=language)  # type: ignore[attr-defined]
                    except Exception:
                        return pytesseract.image_to_string(image)  # type: ignore[attr-defined]
            image = source.convert("RGB")
            try:
                return pytesseract.image_to_string(image, lang=language)  # type: ignore[attr-defined]
            except Exception:
                return pytesseract.image_to_string(image)  # type: ignore[attr-defined]

        return pytesseract_backend, "pytesseract"

    try:
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
    except Exception:
        RapidOCR = None

    if RapidOCR is not None:
        engine = RapidOCR()

        def rapidocr_backend(source: Path | Image.Image) -> str:
            if isinstance(source, Path):
                result, _ = engine(str(source))
            else:
                image = source.convert("RGB")
                result, _ = engine(np.array(image))
            if not result:
                return ""
            texts: list[str] = []
            for item in result:
                if isinstance(item, (list, tuple)) and len(item) >= 2:
                    texts.append(str(item[1]))
                elif isinstance(item, dict) and item.get("text"):
                    texts.append(str(item["text"]))
            return "\n".join(texts)

        return rapidocr_backend, "rapidocr_onnxruntime"

    return None, "unavailable"


def parse_manifest_rows(
    row: Mapping[str, str],
    config: ParseConfig,
    ocr_backend: OcrBackend | None = None,
) -> list[ParsedDocument]:
    dataset = row.get("dataset") or "unknown"
    suffix = (row.get("suffix") or Path(row["file_path"]).suffix).lower()
    if dataset == "dataset3" and suffix == ".xlsx":
        return parse_dataset3_xlsx_rows(row, config)
    return [parse_manifest_row(row, config, ocr_backend)]


def parse_manifest_row(
    row: Mapping[str, str],
    config: ParseConfig,
    ocr_backend: OcrBackend | None = None,
) -> ParsedDocument:
    file_path = Path(row["file_path"])
    dataset = row.get("dataset") or "unknown"
    file_name = row.get("file_name") or file_path.name
    suffix = (row.get("suffix") or file_path.suffix).lower()
    size_bytes = int(row.get("size_bytes") or file_path.stat().st_size)
    doc_id = f"{dataset}:{file_path.stem}"

    title = ""
    raw_text = ""
    clean_text = ""
    table_text = ""
    ocr_text = ""
    page_count: int | None = None
    sheet_count: int | None = None
    image_width: int | None = None
    image_height: int | None = None
    notes: list[str] = []
    parse_status = "ok"

    try:
        if suffix == ".txt":
            raw_text, note = load_text_file(file_path)
            notes.append(note)

        elif suffix == ".docx":
            title, text, table_text, note = extract_docx_text(
                file_path,
                include_ocr=config.include_ocr,
                ocr_backend=ocr_backend,
                ocr_min_text_chars=config.docx_ocr_min_text_chars,
            )
            raw_text = "\n\n".join(part for part in (text, table_text) if part)
            notes.append(note)

        elif suffix == ".pdf":
            title, raw_text, table_text, ocr_text, page_count, note = extract_pdf_text(
                file_path,
                config.max_pdf_pages,
                include_ocr=config.include_ocr,
                ocr_backend=ocr_backend,
                ocr_scale=config.pdf_ocr_scale,
                ocr_min_text_chars=config.pdf_ocr_min_text_chars,
            )
            notes.append(note)
            if ocr_text.strip():
                notes.append("pdf_ocr_used")

        elif suffix == ".xlsx":
            title, raw_text, table_text, sheet_count, note = extract_xlsx_text(
                file_path,
                config.max_xlsx_rows,
                config.max_xlsx_cols,
                include_hidden_sheets=config.include_hidden_sheets,
                include_cell_comments=config.include_cell_comments,
            )
            notes.append(note)

        elif suffix in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}:
            with Image.open(file_path) as image:
                image_width, image_height = image.size
            if config.include_ocr and ocr_backend is not None:
                try:
                    ocr_text = _ocr_source(ocr_backend, file_path)
                    if not ocr_text.strip():
                        notes.append("ocr_empty")
                        parse_status = "needs_ocr"
                    else:
                        notes.append("ocr_used")
                except Exception as exc:
                    notes.append(f"ocr_error={exc}")
                    parse_status = "needs_ocr"
            else:
                parse_status = "needs_ocr"
                notes.append("ocr_unavailable")
            raw_text = ocr_text
            table_text = ""

        else:
            raw_text, note = load_text_file(file_path)
            notes.append(f"fallback_text; {note}")
            parse_status = "unsupported"

    except Exception as exc:
        raw_text = ""
        clean_text = ""
        parse_status = "error"
        notes.append(f"exception={type(exc).__name__}: {exc}")

    combined = "\n\n".join(part for part in (raw_text, table_text, ocr_text) if part)
    combined = normalize_text(combined)
    if not title:
        title = guess_title(combined, file_name)
    time_info = extract_time_info(combined)

    clean_text = combined
    raw_text = combined
    if len(clean_text) > config.max_chars:
        clean_text, truncated = truncate_text(clean_text, config.max_chars)
    else:
        truncated = False
    if len(raw_text) > config.max_chars:
        raw_text, raw_truncated = truncate_text(raw_text, config.max_chars)
        truncated = truncated or raw_truncated

    if parse_status == "ok" and suffix == ".pdf" and not clean_text:
        parse_status = "needs_ocr"
    elif parse_status == "ok" and not clean_text:
        parse_status = "empty"
    if parse_status == "ok" and suffix in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"} and not ocr_text:
        parse_status = "needs_ocr"
    if truncated:
        notes.append(f"truncated_to={config.max_chars}")

    title = strip_surrogates(title)
    time_info = strip_surrogates(time_info)
    raw_text = strip_surrogates(raw_text)
    clean_text = strip_surrogates(clean_text)
    table_text = strip_surrogates(table_text)
    ocr_text = strip_surrogates(ocr_text)
    parse_notes = strip_surrogates("; ".join(note for note in notes if note))

    return ParsedDocument(
        doc_id=doc_id,
        dataset=dataset,
        file_path=str(file_path),
        file_name=file_name,
        suffix=suffix,
        size_bytes=size_bytes,
        title=title,
        time_info=time_info,
        raw_text=raw_text,
        clean_text=clean_text,
        table_text=table_text,
        ocr_text=ocr_text,
        parse_status=parse_status,
        parse_notes=parse_notes,
        page_count=page_count,
        sheet_count=sheet_count,
        image_width=image_width,
        image_height=image_height,
        content_truncated=truncated,
        char_count=len(clean_text),
    )
